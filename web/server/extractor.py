"""文档内容安全智能提取工具：支持 PDF、Word (.docx)、Excel (.xlsx)、表格 (.csv) 及各类纯文本代码文件。"""

from __future__ import annotations

import csv
import logging
from pathlib import Path

log = logging.getLogger("piweb.extractor")


def extract_document_text(file_path: Path, max_chars: int = 40000) -> str:
    """从不同格式的文档中提取纯文本与结构化 Markdown 内容。"""
    if not file_path.is_file():
        return ""

    ext = file_path.suffix.lower()

    try:
        # 1. PDF 提取
        if ext == ".pdf":
            import pypdf

            reader = pypdf.PdfReader(str(file_path))
            pages_text = []
            cur_chars = 0
            for i, page in enumerate(reader.pages):
                txt = page.extract_text() or ""
                if txt.strip():
                    pages_text.append(f"--- [第 {i + 1} 页] ---\n{txt.strip()}")
                    cur_chars += len(txt)
                if cur_chars >= max_chars:
                    pages_text.append(f"\n[提示: 文档篇幅较长，已截取前 {i + 1} 页文字；完整文件已保存在当前工作区]")
                    break
            return "\n\n".join(pages_text)

        # 2. Word (.docx) 提取
        if ext == ".docx":
            import docx

            doc = docx.Document(str(file_path))
            parts = []
            cur_chars = 0

            # 提取段落
            for p in doc.paragraphs:
                txt = p.text.strip()
                if txt:
                    parts.append(txt)
                    cur_chars += len(txt)
                if cur_chars >= max_chars:
                    parts.append("\n[提示: 文档篇幅较长，已截取前部分段落]")
                    break

            # 提取表格
            if cur_chars < max_chars:
                for table_idx, table in enumerate(doc.tables):
                    table_rows = []
                    for row in table.rows:
                        row_cells = [c.text.replace("\n", " ").strip() for c in row.cells]
                        table_rows.append("| " + " | ".join(row_cells) + " |")
                    if table_rows:
                        if len(table_rows) > 1:
                            header_sep = "| " + " | ".join(["---"] * len(table.rows[0].cells)) + " |"
                            table_rows.insert(1, header_sep)
                        parts.append(f"\n[表格 {table_idx + 1}]:\n" + "\n".join(table_rows))
                        cur_chars += sum(len(r) for r in table_rows)
                    if cur_chars >= max_chars:
                        break

            return "\n\n".join(parts)

        # 3. Excel (.xlsx) 提取
        if ext in (".xlsx", ".xlsm"):
            import openpyxl

            wb = openpyxl.load_workbook(str(file_path), data_only=True, read_only=True)
            sheet_parts = []
            cur_chars = 0

            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                rows = []
                # 读取前 100 行
                for r_idx, row in enumerate(ws.iter_rows(values_only=True)):
                    if r_idx >= 100:
                        rows.append("| ... (本工作表仅预览前 100 行数据) |")
                        break
                    row_strs = [str(v) if v is not None else "" for v in row]
                    # 如果整行都是空的则跳过
                    if not any(row_strs):
                        continue
                    rows.append("| " + " | ".join(row_strs) + " |")

                if rows:
                    if len(rows) > 1:
                        # 插入表头分隔
                        col_count = rows[0].count("|") - 1
                        header_sep = "| " + " | ".join(["---"] * max(1, col_count)) + " |"
                        rows.insert(1, header_sep)
                    sheet_parts.append(f"### 工作表 (Sheet): {sheet_name}\n" + "\n".join(rows))
                    cur_chars += sum(len(r) for r in rows)

                if cur_chars >= max_chars:
                    sheet_parts.append("\n[提示: 表格篇幅较长，已截取部分工作表数据]")
                    break

            wb.close()
            return "\n\n".join(sheet_parts)

        # 4. CSV / TSV 提取
        if ext in (".csv", ".tsv"):
            delimiter = "\t" if ext == ".tsv" else ","
            rows = []
            cur_chars = 0

            # 尝试 utf-8 或 gbk 打开
            encodings = ["utf-8", "gbk", "utf-16"]
            for enc in encodings:
                try:
                    with open(file_path, "r", encoding=enc, errors="replace") as f:
                        reader = csv.reader(f, delimiter=delimiter)
                        for r_idx, row in enumerate(reader):
                            if r_idx >= 100:
                                rows.append("| ... (表格数据较多，仅预览前 100 行) |")
                                break
                            rows.append("| " + " | ".join(row) + " |")
                            cur_chars += sum(len(c) for c in row)
                            if cur_chars >= max_chars:
                                break
                    break
                except Exception:
                    continue

            if rows and len(rows) > 1:
                col_count = rows[0].count("|") - 1
                header_sep = "| " + " | ".join(["---"] * max(1, col_count)) + " |"
                rows.insert(1, header_sep)

            return "\n".join(rows)

        # 5. 代码、JSON、Markdown 与纯文本文件
        text_extensions = {
            ".txt", ".md", ".json", ".yaml", ".yml", ".py", ".js", ".ts", ".tsx",
            ".html", ".css", ".sql", ".sh", ".ps1", ".bat", ".xml", ".log", ".ini", ".env"
        }
        if ext in text_extensions or ext == "":
            for enc in ["utf-8", "gbk", "latin-1"]:
                try:
                    content = file_path.read_text(encoding=enc)
                    if len(content) > max_chars:
                        return content[:max_chars] + f"\n\n... (篇幅较长，已截取前 {max_chars} 字符)"
                    return content
                except Exception:
                    continue

    except Exception as exc:
        log.warning("提取文档内容失败 [%s]: %s", file_path.name, exc)
        return f"[提取文档文字时遇到异常: {exc}，完整文件已存入工作区]"

    return ""

